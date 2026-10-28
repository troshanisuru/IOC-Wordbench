"""Writes the results to an Excel file."""

from __future__ import annotations

import collections
from dataclasses import dataclass
import os
from pathlib import Path
import re
from typing import Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

RESULTS_SHEET = "Results"
SUMMARY_SHEET = "Summary"
STATUS_DETECTED = "Detected"
STATUS_CLEAN = "Clean"
STATUS_NO_RESULT = "No result"
STATUS_UNKNOWN = "Unknown"
STATUS_INVALID = "Invalid"
STATUS_ERROR = "Error"
STATUS_NOT_CHECKED = "Not checked"
STATUS_ORDER = (
    STATUS_DETECTED,
    STATUS_CLEAN,
    STATUS_UNKNOWN,
    STATUS_INVALID,
    STATUS_ERROR,
    STATUS_NOT_CHECKED,
)
STATUS_FILLS = {
    "Detected": "F8CBAD",
    "Clean": "C6EFCE",
    "No result": "FFEB9C",
    "Unknown": "FFEB9C",
    "Invalid": "D9D9D9",
    "Error": "D9D9D9",
    "Not checked": "D9D9D9",
}
HEADER_FILL = "D9E1F2"
COLUMN_WIDTHS = (60, 18, 14)
MAX_CELL_CHARS = 32767

_CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class ReportWriteError(Exception):
    """Raised when the report file cannot be written."""


@dataclass(frozen=True)
class ReportRow:
    """Represents a single indicator row in the Excel report."""

    ioc: str
    ioc_type: str = ""
    vt_score: str = ""
    vt_malicious: Optional[int] = None
    vendor_status: str = ""
    vt_last_analysis: str = ""
    checked_at: str = ""
    note: str = ""


@dataclass(frozen=True)
class ReportMeta:
    """Metadata describing the report execution and environment."""

    vendor: str
    input_name: str
    generated_at: str
    duplicates_removed: int = 0


@dataclass(frozen=True)
class ReportInfo:
    """Summary metadata read back from an existing report file."""

    vendor: str
    input_name: str


def _clean_text(text: str) -> str:
    """Removes Excel-incompatible control characters and truncates text."""
    return _CONTROL_CHAR_RE.sub("", text)[:MAX_CELL_CHARS]


def _set_text(cell, text: str) -> None:
    """Sanitizes text and assigns it as a plain-text string cell."""
    text = _clean_text(text)
    if text == "":
        return
    cell.value = text
    cell.data_type = "s"
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        cell.quotePrefix = True


def _headers(vendor: str) -> list[str]:
    """Generates the column header list for the Results sheet."""
    return [
        "IOC",
        f"{vendor} Status",
        "VT Score",
    ]


def write_report(
    path: str | Path, rows: Sequence[ReportRow], meta: ReportMeta
) -> Path:
    """Writes the Results and Summary sheets into an Excel workbook."""
    p = Path(path)
    if p.suffix.lower() != ".xlsx":
        raise ValueError("Report file must end with .xlsx")

    for row in rows:
        if row.vendor_status not in STATUS_FILLS:
            raise ValueError(f"Unsupported status label '{row.vendor_status}'.")

    wb = Workbook()
    ws = wb.active
    ws.title = RESULTS_SHEET

    header_font = Font(bold=True)
    header_fill = PatternFill(
        fill_type="solid", start_color=HEADER_FILL, end_color=HEADER_FILL
    )

    headers = _headers(meta.vendor)
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        _set_text(cell, header_text)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, r in enumerate(rows, start=2):
        cell_a = ws.cell(row=row_idx, column=1)
        _set_text(cell_a, r.ioc)

        cell_b = ws.cell(row=row_idx, column=2)
        _set_text(cell_b, r.vendor_status)
        fill_color = STATUS_FILLS[r.vendor_status]
        cell_b.fill = PatternFill(
            fill_type="solid", start_color=fill_color, end_color=fill_color
        )

        cell_c = ws.cell(row=row_idx, column=3)
        _set_text(cell_c, r.vt_score)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(rows) + 1}"

    col_letters = ("A", "B", "C")
    for letter, width in zip(col_letters, COLUMN_WIDTHS):
        ws.column_dimensions[letter].width = width

    summary_ws = wb.create_sheet(title=SUMMARY_SHEET)
    _set_text(summary_ws["A1"], "Vendor checked")
    _set_text(summary_ws["B1"], meta.vendor)

    _set_text(summary_ws["A2"], "Input file")
    _set_text(summary_ws["B2"], meta.input_name)

    _set_text(summary_ws["A3"], "Generated (UTC)")
    _set_text(summary_ws["B3"], meta.generated_at)

    _set_text(summary_ws["A4"], "Rows in report")
    summary_ws["B4"].value = len(rows)

    _set_text(summary_ws["A5"], "Duplicates removed")
    summary_ws["B5"].value = int(meta.duplicates_removed)

    _set_text(summary_ws["A7"], "Status")
    _set_text(summary_ws["B7"], "Count")
    summary_ws["A7"].font = header_font
    summary_ws["B7"].font = header_font

    counts = collections.Counter(r.vendor_status for r in rows)
    for idx, status in enumerate(STATUS_ORDER, start=8):
        _set_text(summary_ws[f"A{idx}"], status)
        summary_ws[f"B{idx}"].value = counts[status]

    summary_ws.column_dimensions["A"].width = 22
    summary_ws.column_dimensions["B"].width = 40

    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_name(f"{p.stem}.partial{p.suffix}")
    try:
        wb.save(tmp)
        os.replace(tmp, p)
    except PermissionError:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise ReportWriteError(
            f"Cannot write the report. If it is open in Excel, close it and run again: {p}"
        ) from None
    except OSError as exc:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise ReportWriteError(
            f"Cannot write the report ({type(exc).__name__}): {p}"
        ) from None
    except BaseException:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise

    return p


def read_report_info(path: str | Path) -> ReportInfo:
    """Reads the vendor and input file name from the Summary sheet of a report."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"Report file not found: {p}")

    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except KeyError:
        raise ValueError("File is not a vt-ioc-checker report.") from None
    try:
        if (
            RESULTS_SHEET not in wb.sheetnames
            or SUMMARY_SHEET not in wb.sheetnames
        ):
            raise ValueError("File is not a vt-ioc-checker report.")
        ws = wb[SUMMARY_SHEET]
        vals = [
            "" if not row or row[0] is None else str(row[0])
            for row in ws.iter_rows(
                min_row=1, max_row=2, min_col=2, max_col=2, values_only=True
            )
        ]
        while len(vals) < 2:
            vals.append("")
        return ReportInfo(vendor=vals[0], input_name=vals[1])
    finally:
        wb.close()


def read_report_rows(path: str | Path) -> list[ReportRow]:
    """Reads ReportRow entries from the Results sheet of an Excel report."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"Report file not found: {p}")

    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except KeyError:
        raise ValueError("File is not a vt-ioc-checker report.") from None
    try:
        if RESULTS_SHEET not in wb.sheetnames:
            raise ValueError("File is not a vt-ioc-checker report.")
        ws = wb[RESULTS_SHEET]
        rows_iter = ws.iter_rows(min_row=1, max_col=3, values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise ValueError("File is not a vt-ioc-checker report.")

        header = list(header_row) if header_row else []
        while len(header) < 3:
            header.append(None)
        if header[0] != "IOC" or header[2] != "VT Score":
            raise ValueError("File is not a vt-ioc-checker report.")

        rows: list[ReportRow] = []
        for row in rows_iter:
            row_items = list(row) if row else []
            while len(row_items) < 3:
                row_items.append(None)
            if all(c is None for c in row_items[:3]):
                continue
            rows.append(
                ReportRow(
                    ioc="" if row_items[0] is None else str(row_items[0]),
                    vendor_status=(
                        "" if row_items[1] is None else str(row_items[1])
                    ),
                    vt_score="" if row_items[2] is None else str(row_items[2]),
                )
            )
        return rows
    finally:
        wb.close()
