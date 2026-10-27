"""Unit tests for Task 9B small parser, retry, and spreadsheet hardening fixes."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from vt_ioc_checker.ioc_parser import (
    IocType,
    detect_ioc_type,
    load_iocs,
    make_entry,
)
from vt_ioc_checker.report import (
    RESULTS_SHEET,
    ReportMeta,
    ReportRow,
    write_report,
)
from vt_ioc_checker.vt_client import (
    DEFAULT_RATE_LIMIT_WAIT_SECONDS,
    _retry_after_seconds,
)

META = ReportMeta(
    vendor="Sophos",
    input_name="list.txt",
    generated_at="2026-10-25 10:00:00 UTC",
    duplicates_removed=2,
)


def make_row(**overrides) -> ReportRow:
    """Create a ReportRow with configurable field overrides."""
    defaults = {
        "ioc": "evil.com",
        "ioc_type": "Domain",
        "vt_score": "5/66",
        "vt_malicious": 5,
        "vendor_status": "Detected",
        "vt_last_analysis": "2023-11-14 22:13:20 UTC",
        "checked_at": "2026-10-25 10:00:00 UTC",
        "note": "",
    }
    defaults.update(overrides)
    return ReportRow(**defaults)


@pytest.mark.parametrize("value", ["nan", "inf", "-inf"])
def test_retry_after_non_finite(value: str) -> None:
    """Return the default wait duration when Retry-After is non-finite."""
    resp = SimpleNamespace(headers={"Retry-After": value})
    assert _retry_after_seconds(resp) == DEFAULT_RATE_LIMIT_WAIT_SECONDS


@pytest.mark.parametrize("header", ['"ioc"', "'Indicator'"])
def test_quoted_header_skipped(tmp_path: Path, header: str) -> None:
    """Skip a first line that is a quoted header name in a text file."""
    path = tmp_path / "list.txt"
    path.write_text(f"{header}\n8.8.8.8\n", encoding="utf-8")
    res = load_iocs(path)
    assert res.total_read == 1
    assert [e.value for e in res.entries] == ["8.8.8.8"]


def test_too_long_value_invalid() -> None:
    """Accept an 8000-char URL and mark an 8001-char URL as Invalid."""
    base = "http://evil.com/"
    v8000 = base + "a" * (8000 - len(base))
    assert detect_ioc_type(v8000) == IocType.URL
    assert detect_ioc_type(v8000 + "a") == IocType.INVALID


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("http://1.2.3.4:99999/x", IocType.INVALID),
        ("http://1.2.3.4:0/x", IocType.INVALID),
        ("http://evil.com:abc/x", IocType.INVALID),
        ("http://evil..com/x", IocType.INVALID),
        ("http://-bad.com/x", IocType.INVALID),
        ("http://evil.com:8080/x", IocType.URL),
        ("evil.com:99999/x", IocType.INVALID),
        ("evil.com:0/x", IocType.INVALID),
        ("evil.com:65535/x", IocType.URL),
    ],
)
def test_url_port_and_host_validation(value: str, expected: IocType) -> None:
    """Validate URL ports and hostnames in both scheme and schemeless URLs."""
    assert detect_ioc_type(value) == expected


def test_quote_prefix_for_formula_like_text(tmp_path: Path) -> None:
    """Set quotePrefix on formula-like cell strings while keeping values intact."""
    path = tmp_path / "r.xlsx"
    rows = [
        make_row(ioc="=1+1", vendor_status="Invalid"),
        make_row(ioc="+cmd", vendor_status="Invalid"),
        make_row(ioc="-2+3", vendor_status="Invalid"),
        make_row(ioc="@SUM(A1)", vendor_status="Invalid"),
        make_row(ioc="evil.com", vendor_status="Clean"),
    ]
    write_report(path, rows, META)
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[RESULTS_SHEET]
        expected_vals = ["=1+1", "+cmd", "-2+3", "@SUM(A1)"]
        for idx, expected_val in enumerate(expected_vals, start=2):
            cell = ws[f"A{idx}"]
            assert cell.quotePrefix is True
            assert cell.value == expected_val
        assert not ws["A6"].quotePrefix
        assert ws["A6"].value == "evil.com"
    finally:
        wb.close()


@pytest.mark.parametrize(
    ("raw", "expected_value"),
    [
        ("HXXPS://Evil[.]com/x", "https://Evil.com/x"),
        ("HXXP://a[.]com/x", "http://a.com/x"),
    ],
)
def test_hxxps_scheme_lowercased(raw: str, expected_value: str) -> None:
    """Ensure refanging HXXPS/HXXP produces a lowercase http/https scheme."""
    entry = make_entry(raw)
    assert entry.value == expected_value
    assert entry.ioc_type == IocType.URL
