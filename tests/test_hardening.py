"""Unit tests for Task 9A hardening fixes."""

from __future__ import annotations

from datetime import datetime, timezone
import io
from pathlib import Path
import zipfile

import openpyxl
import pytest
import requests

from vt_ioc_checker.__main__ import main
from vt_ioc_checker.config import ConfigError, load_settings
from vt_ioc_checker.ioc_parser import load_iocs
from vt_ioc_checker.report import (
    ReportInfo,
    ReportMeta,
    ReportWriteError,
    read_report_info,
    read_report_rows,
    write_report,
)
from vt_ioc_checker.runner import run_check
from vt_ioc_checker.vt_client import RateLimiter, VTClient, VTResult

NOW = datetime(2026, 10, 25, 10, 0, 0, tzinfo=timezone.utc)
MALICIOUS = VTResult(
    True, {"malicious": 5, "undetected": 60}, {"Sophos": "malicious"}, 1700000000
)
CLEAN = VTResult(
    True,
    {"malicious": 40, "undetected": 30},
    {"Sophos": "undetected"},
    1700000000,
)
THREE = ["evil.com", "8.8.8.8", "example.com"]
KEY = "test-key-123456"
META = ReportMeta("Sophos", "list.txt", "2026-10-25 10:00:00 UTC")


class FakeClient:
    """Fake VirusTotal client for testing pipeline logic."""

    def __init__(self, results=None) -> None:
        """Initialize the fake client with optional canned results."""
        self.results = results or {}
        self.calls = []

    def lookup(self, ioc_type, value):
        """Record the lookup call and return or raise the configured result."""
        self.calls.append((ioc_type, value))
        item = self.results.get(value, VTResult(found=False))
        if isinstance(item, BaseException):
            raise item
        return item


class FakeResponse:
    """Fake HTTP response for testing VTClient error paths."""

    def __init__(
        self,
        status_code: int = 200,
        json_data: dict | None = None,
        headers: dict[str, str] | None = None,
        bad_json: bool = False,
    ) -> None:
        """Initialize the fake HTTP response."""
        self.status_code = status_code
        self._json_data = json_data
        self.headers = headers if headers is not None else {}
        self.bad_json = bad_json

    def json(self) -> dict:
        """Return the JSON payload or raise ValueError when bad_json is set."""
        if self.bad_json:
            raise ValueError("bad json")
        return self._json_data if self._json_data is not None else {}


class FakeSession:
    """Fake requests.Session queueing responses or exceptions."""

    def __init__(self, responses: list) -> None:
        """Initialize the session with a list of responses or exceptions."""
        self.responses = list(responses)
        self.requests: list[tuple[str, dict | None, float | None]] = []

    def get(
        self,
        url: str,
        headers: dict | None = None,
        timeout: float | None = None,
    ) -> FakeResponse:
        """Record the GET request and return or raise the next queued item."""
        self.requests.append((url, headers, timeout))
        resp = self.responses.pop(0)
        if isinstance(resp, Exception):
            raise resp
        return resp


def write_list(tmp_path: Path, lines, name: str = "in.txt") -> Path:
    """Write an IOC list file in UTF-8 and return its path."""
    path = tmp_path / name
    path.write_bytes(("\n".join(lines) + "\n").encode("utf-8"))
    return path


def run(
    tmp_path: Path,
    client: FakeClient,
    lines=None,
    name: str = "in.txt",
    out_name: str = "out.xlsx",
    **kwargs,
):
    """Run run_check with standard test defaults and return (summary, out)."""
    in_path = write_list(tmp_path, lines or THREE, name=name)
    out = tmp_path / out_name
    kwargs.setdefault("now", lambda: NOW)
    summary = run_check(in_path, out, client, **kwargs)
    return summary, out


def test_resume_rejects_other_vendor(tmp_path: Path) -> None:
    """Reject resuming a report that was generated for a different vendor."""
    client1 = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    _, out = run(tmp_path, client1)
    before = read_report_rows(out)
    client2 = FakeClient()
    with pytest.raises(ValueError, match="made for vendor 'Sophos'"):
        run(tmp_path, client2, vendor="Microsoft")
    assert client2.calls == []
    assert read_report_rows(out) == before


def test_resume_rejects_other_input_name(tmp_path: Path) -> None:
    """Reject resuming a report that was generated from a different input file."""
    client1 = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    _, out = run(tmp_path, client1, name="in.txt")
    before = read_report_rows(out)
    client2 = FakeClient()
    with pytest.raises(ValueError, match="made from input file 'in.txt'"):
        run(tmp_path, client2, name="other.txt")
    assert client2.calls == []
    assert read_report_rows(out) == before


def test_resume_vendor_check_is_case_insensitive(tmp_path: Path) -> None:
    """Allow resuming when the vendor name differs only by letter case."""
    client1 = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    run(tmp_path, client1)
    client2 = FakeClient()
    summary2, _ = run(tmp_path, client2, vendor="sophos")
    assert summary2.reused == 3
    assert client2.calls == []


def test_no_resume_allows_other_vendor(tmp_path: Path) -> None:
    """Allow replacing a report with another vendor when resume is False."""
    run(tmp_path, FakeClient({"evil.com": MALICIOUS}))
    _, out = run(tmp_path, FakeClient(), vendor="Microsoft", resume=False)
    wb = openpyxl.load_workbook(out)
    try:
        assert wb["Results"]["E1"].value == "Microsoft Status"
    finally:
        wb.close()


def test_read_report_info(tmp_path: Path) -> None:
    """Read vendor and input file name from the Summary sheet."""
    path = tmp_path / "r.xlsx"
    write_report(
        path, [], ReportMeta("Sophos", "list.txt", "2026-10-25 10:00:00 UTC")
    )
    assert read_report_info(path) == ReportInfo("Sophos", "list.txt")


def test_read_report_info_missing_file(tmp_path: Path) -> None:
    """Raise FileNotFoundError when the report file does not exist."""
    with pytest.raises(FileNotFoundError):
        read_report_info(tmp_path / "nope.xlsx")


def test_read_report_info_other_workbook(tmp_path: Path) -> None:
    """Raise ValueError when the workbook is not a vt-ioc-checker report."""
    path = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="not a vt-ioc-checker report"):
        read_report_info(path)


def test_final_save_retries_then_succeeds(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Retry the final save on ReportWriteError until it succeeds."""
    from vt_ioc_checker import runner

    real_write = runner.write_report
    calls = [0]

    def flaky_write(path, rows, meta):
        """Raise ReportWriteError on the first two calls then succeed."""
        calls[0] += 1
        if calls[0] <= 2:
            raise ReportWriteError("locked")
        return real_write(path, rows, meta)

    monkeypatch.setattr(runner, "write_report", flaky_write)
    sleeps: list[float] = []
    summary, out = run(
        tmp_path,
        FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN}),
        sleep=sleeps.append,
        final_save_wait=10.0,
    )
    assert sleeps == [10.0, 10.0]
    assert summary.stop_kind == ""
    assert summary.report_path == out
    rows = read_report_rows(out)
    assert len(rows) == 3
    assert all(r.vendor_status != "Not checked" for r in rows)


def test_final_save_falls_back_to_unsaved_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Save to a timestamped fallback file when all final save retries fail."""
    from vt_ioc_checker import runner

    real_write = runner.write_report
    target = tmp_path / "out.xlsx"

    def locked_target(path, rows, meta):
        """Raise ReportWriteError for target path and delegate for fallback."""
        if Path(path) == target:
            raise ReportWriteError("locked")
        return real_write(path, rows, meta)

    monkeypatch.setattr(runner, "write_report", locked_target)
    sleeps: list[float] = []
    logs: list[str] = []
    summary, out = run(
        tmp_path,
        FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN}),
        sleep=sleeps.append,
        final_save_wait=10.0,
        log=logs.append,
    )
    expected = tmp_path / "out.unsaved-20261025-100000.xlsx"
    assert sleeps == [10.0] * 5
    assert summary.report_path == expected
    assert expected.exists()
    rows = read_report_rows(expected)
    assert len(rows) == 3
    assert all(r.vendor_status != "Not checked" for r in rows)
    assert not out.exists()
    assert any("Results were saved to" in msg for msg in logs)


def test_final_save_everything_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Propagate ReportWriteError when both primary and fallback saves fail."""
    from vt_ioc_checker import runner

    def always_locked(path, rows, meta):
        """Always raise ReportWriteError."""
        raise ReportWriteError("locked")

    monkeypatch.setattr(runner, "write_report", always_locked)
    sleeps: list[float] = []
    with pytest.raises(ReportWriteError):
        run(tmp_path, FakeClient(), sleep=sleeps.append, final_save_wait=10.0)
    assert sleeps == [10.0] * 5


def test_ctrl_c_while_waiting_goes_to_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Write immediately to the fallback file if Ctrl+C interrupts retry sleep."""
    from vt_ioc_checker import runner

    real_write = runner.write_report
    target = tmp_path / "out.xlsx"

    def locked_target(path, rows, meta):
        """Raise ReportWriteError for target path and delegate for fallback."""
        if Path(path) == target:
            raise ReportWriteError("locked")
        return real_write(path, rows, meta)

    def interrupt_sleep(seconds: float) -> None:
        """Raise KeyboardInterrupt when retry sleep is called."""
        raise KeyboardInterrupt

    monkeypatch.setattr(runner, "write_report", locked_target)
    summary, _ = run(
        tmp_path, FakeClient(), sleep=interrupt_sleep, final_save_wait=10.0
    )
    expected = tmp_path / "out.unsaved-20261025-100000.xlsx"
    assert summary.report_path == expected
    assert expected.exists()


def test_cli_reports_fallback_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Return EXIT_FILE and report the fallback path when primary save fails."""
    from vt_ioc_checker import runner

    real_write = runner.write_report

    def locked_default(path, rows, meta):
        """Raise ReportWriteError for the default output filename."""
        if Path(path).name == "list_sophos_report.xlsx":
            raise ReportWriteError("locked")
        return real_write(path, rows, meta)

    monkeypatch.setattr(runner, "write_report", locked_default)
    monkeypatch.delenv("VT_API_KEY", raising=False)
    monkeypatch.delenv("VT_REQUESTS_PER_MINUTE", raising=False)
    monkeypatch.chdir(tmp_path)
    src = write_list(tmp_path, THREE, name="list.txt")
    env = tmp_path / ".env"
    env.write_text(f"VT_API_KEY={KEY}\n", encoding="utf-8")
    out, err = io.StringIO(), io.StringIO()
    code = main(
        [str(src), "--env-file", str(env)],
        client_factory=lambda s: FakeClient({"evil.com": MALICIOUS}),
        stdout=out,
        stderr=err,
    )
    assert code == 3
    assert "Results were saved to" in err.getvalue()
    matches = list(
        (tmp_path / "data" / "output").glob(
            "list_sophos_report.unsaved-*.xlsx"
        )
    )
    assert len(matches) == 1


def test_write_report_cleans_tmp_on_keyboard_interrupt(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Remove the temporary .partial file if KeyboardInterrupt occurs in write_report."""
    def interrupt_replace(src, dst):
        """Raise KeyboardInterrupt during os.replace."""
        raise KeyboardInterrupt

    monkeypatch.setattr("vt_ioc_checker.report.os.replace", interrupt_replace)
    target = tmp_path / "r.xlsx"
    with pytest.raises(KeyboardInterrupt):
        write_report(target, [], META)
    assert list(tmp_path.glob("*.partial*")) == []
    assert not target.exists()


def test_interrupt_during_logging_saves_progress(tmp_path: Path) -> None:
    """Save finished lookups when KeyboardInterrupt happens outside client.lookup."""
    def interrupting_log(msg: str) -> None:
        """Raise KeyboardInterrupt on the second progress log line."""
        if msg.startswith("[2/3]"):
            raise KeyboardInterrupt

    summary, out = run(
        tmp_path,
        FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN}),
        log=interrupting_log,
    )
    assert summary.stop_kind == "interrupted"
    assert [r.vendor_status for r in read_report_rows(out)] == [
        "Detected",
        "Clean",
        "Not checked",
    ]
    assert summary.looked_up == 2


def test_env_file_not_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Raise ConfigError with a UTF-8 message when .env is UTF-16 encoded."""
    monkeypatch.delenv("VT_API_KEY", raising=False)
    monkeypatch.delenv("VT_REQUESTS_PER_MINUTE", raising=False)
    path = tmp_path / ".env"
    path.write_bytes("VT_API_KEY=abc123\n".encode("utf-16"))
    with pytest.raises(ConfigError, match="UTF-8"):
        load_settings(path)


def test_csv_field_too_large(tmp_path: Path) -> None:
    """Raise ValueError when a CSV field exceeds the parser size limit."""
    path = tmp_path / "list.csv"
    path.write_text("a" * 200000 + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="could not be parsed"):
        load_iocs(path)


def test_cli_catch_all(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Return EXIT_ERROR and hide exception details when an unexpected error occurs."""
    monkeypatch.delenv("VT_API_KEY", raising=False)
    monkeypatch.delenv("VT_REQUESTS_PER_MINUTE", raising=False)
    monkeypatch.chdir(tmp_path)
    src = write_list(tmp_path, THREE, name="list.txt")
    env = tmp_path / ".env"
    env.write_text(f"VT_API_KEY={KEY}\n", encoding="utf-8")

    def boom_factory(settings):
        """Raise an unexpected RuntimeError."""
        raise RuntimeError("boom")

    out, err = io.StringIO(), io.StringIO()
    code = main(
        [str(src), "--env-file", str(env)],
        client_factory=boom_factory,
        stdout=out,
        stderr=err,
    )
    assert code == 7
    assert "Unexpected error (RuntimeError)" in err.getvalue()
    assert "boom" not in out.getvalue()
    assert "boom" not in err.getvalue()


def test_suspicious_note(tmp_path: Path) -> None:
    """Add an explanatory note when the vendor verdict category is suspicious."""
    res = VTResult(
        True, {"suspicious": 1, "undetected": 70}, {"Sophos": "suspicious"}, None
    )
    _, out = run(tmp_path, FakeClient({"evil.com": res}), lines=["evil.com"])
    rows = read_report_rows(out)
    assert len(rows) == 1
    assert rows[0].vendor_status == "Detected"
    assert rows[0].vt_score == "0/71"
    assert rows[0].note == "Sophos verdict is suspicious (not malicious)"


@pytest.mark.parametrize(
    ("responses", "expected_exit_code"),
    [
        pytest.param(
            [FakeResponse(401, {"error": {"code": "WrongCredentialsError"}})],
            4,
            id="401",
        ),
        pytest.param(
            [FakeResponse(429, {"error": {"code": "QuotaExceededError"}})],
            5,
            id="quota",
        ),
        pytest.param(
            [
                requests.exceptions.ConnectionError(),
                requests.exceptions.ConnectionError(),
            ],
            0,
            id="network",
        ),
        pytest.param(
            [FakeResponse(500), FakeResponse(500)],
            0,
            id="server",
        ),
        pytest.param(
            [FakeResponse(200, bad_json=True)],
            0,
            id="badjson",
        ),
    ],
)
def test_cli_never_leaks_key_with_real_client(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    responses: list,
    expected_exit_code: int,
) -> None:
    """Ensure the API key never leaks to output or report files across error modes."""
    monkeypatch.delenv("VT_API_KEY", raising=False)
    monkeypatch.delenv("VT_REQUESTS_PER_MINUTE", raising=False)
    monkeypatch.chdir(tmp_path)
    src = write_list(tmp_path, ["evil.com"], name="list.txt")
    env = tmp_path / ".env"
    env.write_text(f"VT_API_KEY={KEY}\n", encoding="utf-8")
    session = FakeSession(responses)
    out, err = io.StringIO(), io.StringIO()
    code = main(
        [str(src), "--env-file", str(env)],
        client_factory=lambda s: VTClient(
            s.api_key,
            RateLimiter(4, clock=lambda: 0.0, sleep=lambda x: None),
            session=session,
            sleep=lambda x: None,
            max_retries=1,
        ),
        stdout=out,
        stderr=err,
    )
    assert code == expected_exit_code
    assert session.requests[0][1] is not None
    assert session.requests[0][1]["x-apikey"] == KEY
    assert KEY not in out.getvalue()
    assert KEY not in err.getvalue()
    report_path = tmp_path / "data" / "output" / "list_sophos_report.xlsx"
    if report_path.exists():
        wb = openpyxl.load_workbook(report_path)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cell_val in row:
                        if cell_val is not None:
                            assert KEY not in str(cell_val)
        finally:
            wb.close()
        with zipfile.ZipFile(report_path, "r") as zf:
            for member in zf.namelist():
                text = zf.read(member).decode("utf-8", errors="replace")
                assert KEY not in text
