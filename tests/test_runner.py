"""Unit tests for the pipeline runner."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from vt_ioc_checker.ioc_parser import IocType
from vt_ioc_checker.report import (
    ReportRow,
    ReportWriteError,
    read_report_rows,
)
from vt_ioc_checker.runner import (
    DEFAULT_SAVE_EVERY,
    run_check,
)
from vt_ioc_checker.vt_client import (
    VTAuthError,
    VTQuotaExceededError,
    VTRequestError,
    VTResult,
)

NOW = datetime(2026, 10, 25, 10, 0, 0, tzinfo=timezone.utc)
CHECKED = "2026-10-25 10:00:00 UTC"
NOW2 = datetime(2026, 10, 26, 9, 0, 0, tzinfo=timezone.utc)
CHECKED2 = "2026-10-26 09:00:00 UTC"
MD5 = "44d88612fea8a8f36de82e1278abb02f"
MALICIOUS = VTResult(
    True, {"malicious": 5, "undetected": 60}, {"Sophos": "malicious"}, 1700000000
)
CLEAN = VTResult(
    True,
    {"malicious": 40, "undetected": 30},
    {"Sophos": "undetected"},
    1700000000,
)


class FakeClient:
    """Fake VirusTotal client for testing pipeline logic."""

    def __init__(self, results=None):
        self.results = results or {}
        self.calls = []

    def lookup(self, ioc_type, value):
        self.calls.append((ioc_type, value))
        item = self.results.get(value, VTResult(found=False))
        if isinstance(item, BaseException):
            raise item
        return item


BASIC_LINES = ["evil.com", "8.8.8.8", "not an ioc", "EVIL.com", MD5]
THREE = ["evil.com", "8.8.8.8", "example.com"]


def run(tmp_path: Path, client: FakeClient, lines=None, **kwargs):
    """Helper to run run_check with standard defaults."""
    in_path = tmp_path / "in.txt"
    in_path.write_bytes(
        ("\n".join(lines or BASIC_LINES) + "\n").encode("utf-8")
    )
    out = tmp_path / "out.xlsx"
    kwargs.setdefault("now", lambda: NOW)
    summary = run_check(in_path, out, client, **kwargs)
    return summary, out


def test_basic_run(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    summary, out = run(tmp_path, client)
    assert summary.total_unique == 4
    assert summary.duplicates_removed == 1
    assert summary.invalid == 1
    assert summary.reused == 0
    assert summary.looked_up == 3
    assert summary.stop_kind == ""
    assert summary.stopped_early is False
    assert summary.report_path == out
    assert summary.counts == {
        "Detected": 1,
        "Clean": 1,
        "No result": 0,
        "Unknown": 1,
        "Invalid": 1,
        "Error": 0,
        "Not checked": 0,
    }
    assert client.calls == [
        (IocType.DOMAIN, "evil.com"),
        (IocType.IP, "8.8.8.8"),
        (IocType.MD5, MD5),
    ]
    assert read_report_rows(out) == [
        ReportRow(
            "evil.com",
            "Domain",
            "5/65",
            5,
            "Detected",
            "2023-11-14 22:13:20 UTC",
            CHECKED,
            "",
        ),
        ReportRow(
            "8.8.8.8",
            "IP",
            "40/70",
            40,
            "Clean",
            "2023-11-14 22:13:20 UTC",
            CHECKED,
            "",
        ),
        ReportRow(
            "not an ioc",
            "Invalid",
            "",
            None,
            "Invalid",
            "",
            "",
            "Not a valid IOC; not looked up",
        ),
        ReportRow(
            MD5,
            "MD5",
            "Unknown",
            None,
            "Unknown",
            "",
            CHECKED,
            "Not found in VirusTotal",
        ),
    ]


@pytest.mark.parametrize(
    "result,expected_note",
    [
        (
            VTResult(
                True,
                {"malicious": 2, "undetected": 60},
                {"Microsoft": "malicious"},
                None,
            ),
            "Sophos gave no verdict",
        ),
        (
            VTResult(True, {"undetected": 10}, {"Sophos": "timeout"}, None),
            "Sophos gave no verdict (timeout)",
        ),
    ],
)
def test_no_result_note(
    tmp_path: Path, result: VTResult, expected_note: str
) -> None:
    client = FakeClient({"evil.com": result})
    summary, out = run(tmp_path, client, lines=["evil.com"])
    rows = read_report_rows(out)
    assert len(rows) == 1
    assert rows[0].vendor_status == "No result"
    assert rows[0].note == expected_note


def test_other_vendor(tmp_path: Path) -> None:
    result = VTResult(
        True,
        {"malicious": 2, "undetected": 60},
        {"Microsoft": "malicious"},
        None,
    )
    client = FakeClient({"evil.com": result})
    summary, out = run(tmp_path, client, lines=["evil.com"], vendor="Microsoft")
    rows = read_report_rows(out)
    assert rows[0].vendor_status == "Detected"
    wb = openpyxl.load_workbook(out)
    assert wb["Results"]["E1"].value == "Microsoft Status"
    wb.close()


def test_request_error_not_fatal(tmp_path: Path) -> None:
    client = FakeClient(
        {"evil.com": VTRequestError("request failed"), "8.8.8.8": CLEAN}
    )
    summary, out = run(tmp_path, client, lines=THREE)
    rows = read_report_rows(out)
    assert rows[0].vendor_status == "Error"
    assert rows[0].vt_score == ""
    assert rows[0].vt_malicious is None
    assert rows[0].checked_at == CHECKED
    assert rows[0].note == "request failed"
    assert rows[1].vendor_status == "Clean"
    assert rows[2].vendor_status == "Unknown"
    assert summary.stop_kind == ""
    assert summary.counts["Error"] == 1
    assert summary.looked_up == 3


def test_auth_error_stops_run(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": VTAuthError("auth failed")})
    summary, out = run(tmp_path, client, lines=THREE)
    assert summary.stop_kind == "auth"
    assert summary.stop_reason == "auth failed"
    assert summary.looked_up == 1
    assert len(client.calls) == 1
    rows = read_report_rows(out)
    assert [r.vendor_status for r in rows] == [
        "Not checked",
        "Not checked",
        "Not checked",
    ]
    assert rows[0].note == "auth failed"
    assert rows[1].note == "Run stopped before this IOC was checked"
    assert rows[2].note == "Run stopped before this IOC was checked"
    assert out.exists()


def test_quota_error_stops_run(tmp_path: Path) -> None:
    client = FakeClient(
        {"evil.com": MALICIOUS, "8.8.8.8": VTQuotaExceededError("quota hit")}
    )
    summary, out = run(tmp_path, client, lines=THREE)
    assert summary.stop_kind == "quota"
    assert summary.looked_up == 2
    rows = read_report_rows(out)
    assert [r.vendor_status for r in rows] == [
        "Detected",
        "Not checked",
        "Not checked",
    ]
    assert rows[1].note == "quota hit"
    assert rows[2].note == "Run stopped before this IOC was checked"


def test_max_lookups(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    summary, out = run(tmp_path, client, lines=THREE, max_lookups=2)
    assert summary.stop_kind == "limit"
    assert summary.stop_reason == "Reached the lookup limit for this run (2)."
    assert summary.looked_up == 2
    rows = read_report_rows(out)
    assert [r.vendor_status for r in rows] == [
        "Detected",
        "Clean",
        "Not checked",
    ]
    assert rows[2].note == "Reached the lookup limit for this run (2)."


def test_interrupt_saves_progress(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": KeyboardInterrupt()})
    summary, out = run(tmp_path, client, lines=THREE)
    assert summary.stop_kind == "interrupted"
    assert summary.stop_reason == "Interrupted by user (Ctrl+C)."
    assert summary.looked_up == 2
    rows = read_report_rows(out)
    assert [r.vendor_status for r in rows] == [
        "Detected",
        "Not checked",
        "Not checked",
    ]
    assert rows[1].note == "Interrupted by user (Ctrl+C)."
    assert rows[2].note == "Interrupted by user (Ctrl+C)."


def test_unexpected_error_saves_progress(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": RuntimeError("boom")})
    summary, out = run(tmp_path, client, lines=THREE)
    assert summary.stop_kind == "error"
    assert summary.stop_reason == "Unexpected error (RuntimeError)."
    assert "boom" not in summary.stop_reason
    rows = read_report_rows(out)
    assert [r.vendor_status for r in rows] == [
        "Detected",
        "Not checked",
        "Not checked",
    ]
    assert rows[1].note == "Unexpected error (RuntimeError)."
    assert rows[2].note == "Run stopped before this IOC was checked"


def test_resume_reuses_final_rows(tmp_path: Path) -> None:
    client1 = FakeClient({"evil.com": MALICIOUS})
    summary1, out = run(tmp_path, client1, lines=THREE, max_lookups=1)
    client2 = FakeClient({"8.8.8.8": CLEAN})
    summary2, out = run(tmp_path, client2, lines=THREE, now=lambda: NOW2)
    assert client2.calls == [
        (IocType.IP, "8.8.8.8"),
        (IocType.DOMAIN, "example.com"),
    ]
    assert summary2.reused == 1
    assert summary2.looked_up == 2
    assert summary2.stop_kind == ""
    rows = read_report_rows(out)
    assert rows[0].checked_at == CHECKED
    assert rows[1].checked_at == CHECKED2
    assert rows[2].checked_at == CHECKED2
    assert rows[0].vendor_status == "Detected"


def test_resume_rechecks_error_and_not_checked(tmp_path: Path) -> None:
    client1 = FakeClient({
        "evil.com": VTRequestError("request failed"),
        "8.8.8.8": CLEAN,
        "example.com": VTQuotaExceededError("quota hit"),
    })
    summary1, out = run(tmp_path, client1, lines=THREE)
    rows1 = read_report_rows(out)
    assert [r.vendor_status for r in rows1] == ["Error", "Clean", "Not checked"]

    client2 = FakeClient({"evil.com": MALICIOUS, "example.com": CLEAN})
    summary2, out = run(tmp_path, client2, lines=THREE)
    assert client2.calls == [
        (IocType.DOMAIN, "evil.com"),
        (IocType.DOMAIN, "example.com"),
    ]
    assert summary2.reused == 1
    rows2 = read_report_rows(out)
    assert [r.vendor_status for r in rows2] == ["Detected", "Clean", "Clean"]


def test_resume_false_ignores_previous(tmp_path: Path) -> None:
    client1 = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    summary1, out = run(tmp_path, client1, lines=THREE)
    client2 = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    summary2, out = run(tmp_path, client2, lines=THREE, resume=False)
    assert len(client2.calls) == 3
    assert summary2.reused == 0


def test_resume_drops_iocs_not_in_input(tmp_path: Path) -> None:
    client1 = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    summary1, out = run(tmp_path, client1, lines=THREE)
    client2 = FakeClient()
    lines2 = ["evil.com", "8.8.8.8", "evil.org"]
    summary2, out = run(tmp_path, client2, lines=lines2)
    assert client2.calls == [(IocType.DOMAIN, "evil.org")]
    assert summary2.reused == 2
    rows = read_report_rows(out)
    assert [r.ioc for r in rows] == ["evil.com", "8.8.8.8", "evil.org"]


def test_non_report_file_not_overwritten(tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    wb = openpyxl.Workbook()
    wb.save(out)
    wb.close()
    in_path = tmp_path / "in.txt"
    in_path.write_bytes(b"evil.com\n")
    client = FakeClient()
    with pytest.raises(ValueError, match="not a vt-ioc-checker report"):
        run_check(in_path, out, client)
    assert client.calls == []
    wb2 = openpyxl.load_workbook(out)
    assert wb2.sheetnames == ["Sheet"]
    wb2.close()


def test_periodic_save(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from vt_ioc_checker import runner

    real_write_report = runner.write_report
    not_checked_counts = []

    def spy_write_report(path, rows, meta):
        not_checked_counts.append(
            sum(1 for r in rows if r.vendor_status == "Not checked")
        )
        return real_write_report(path, rows, meta)

    monkeypatch.setattr(runner, "write_report", spy_write_report)
    lines = ["evil.com", "8.8.8.8", "example.com", "test.org", "sample.net"]
    client = FakeClient()
    summary, out = run(tmp_path, client, lines=lines, save_every=2)
    assert not_checked_counts == [3, 1, 0]


def test_periodic_save_failure_does_not_stop_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from vt_ioc_checker import runner

    real_write_report = runner.write_report
    calls = [0]

    def failing_write_report(path, rows, meta):
        calls[0] += 1
        if calls[0] <= 2:
            raise ReportWriteError("locked")
        return real_write_report(path, rows, meta)

    monkeypatch.setattr(runner, "write_report", failing_write_report)
    log_messages = []
    summary, out = run(
        tmp_path,
        FakeClient(),
        lines=THREE,
        save_every=1,
        log=log_messages.append,
    )
    assert summary.stop_kind == ""
    rows = read_report_rows(out)
    assert all(r.vendor_status != "Not checked" for r in rows)
    assert any("Could not save progress" in msg for msg in log_messages)


def test_final_save_failure_propagates(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from vt_ioc_checker import runner

    def fail_write(path, rows, meta):
        raise ReportWriteError("disk full")

    monkeypatch.setattr(runner, "write_report", fail_write)
    with pytest.raises(ReportWriteError):
        run(tmp_path, FakeClient(), lines=THREE)


def test_progress_log(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    messages = []
    summary, out = run(tmp_path, client, lines=THREE, log=messages.append)
    assert len(messages) >= 3
    assert messages[0].startswith("[1/3] evil.com")
    assert all(m.isascii() for m in messages)


@pytest.mark.parametrize(
    "arg,value",
    [
        ("max_lookups", 0),
        ("save_every", 0),
        ("vendor", "  "),
    ],
)
def test_argument_validation(tmp_path: Path, arg: str, value: Any) -> None:
    in_path = tmp_path / "in.txt"
    in_path.write_text("evil.com\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    client = FakeClient()
    with pytest.raises(ValueError):
        run_check(in_path, out, client, **{arg: value})
    assert client.calls == []


def test_missing_input(tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    client = FakeClient()
    with pytest.raises(FileNotFoundError):
        run_check(tmp_path / "nope.txt", out, client)
    assert not out.exists()


def test_all_invalid_input(tmp_path: Path) -> None:
    client = FakeClient()
    summary, out = run(tmp_path, client, lines=["not an ioc", "=cmd"])
    assert summary.looked_up == 0
    assert client.calls == []
    assert summary.counts["Invalid"] == 2
    rows = read_report_rows(out)
    assert [r.ioc for r in rows] == ["not an ioc", "=cmd"]
    assert out.exists()


def test_report_meta(tmp_path: Path) -> None:
    client = FakeClient({"evil.com": MALICIOUS, "8.8.8.8": CLEAN})
    summary, out = run(tmp_path, client)
    wb = openpyxl.load_workbook(out)
    summary_sheet = wb["Summary"]
    assert summary_sheet["B1"].value == "Sophos"
    assert summary_sheet["B2"].value == "in.txt"
    assert summary_sheet["B3"].value == CHECKED
    assert summary_sheet["B4"].value == 4
    assert summary_sheet["B5"].value == 1
    wb.close()


def test_refanged_value_is_looked_up(tmp_path: Path) -> None:
    client = FakeClient()
    summary, out = run(tmp_path, client, lines=["hxxp://evil[.]com/a"])
    assert client.calls == [(IocType.URL, "http://evil.com/a")]
    rows = read_report_rows(out)
    assert rows[0].ioc == "http://evil.com/a"
    assert rows[0].ioc_type == "URL"
