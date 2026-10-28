"""Unit tests for Excel report generation and parsing."""

from __future__ import annotations

import os
from pathlib import Path
import zipfile

import openpyxl
import pytest

from vt_ioc_checker.report import (
    RESULTS_SHEET,
    SUMMARY_SHEET,
    ReportMeta,
    ReportRow,
    ReportWriteError,
    read_report_rows,
    write_report,
)

META = ReportMeta(
    vendor="Sophos",
    input_name="list.txt",
    generated_at="2026-10-25 10:00:00 UTC",
    duplicates_removed=2,
)


def make_row(**overrides) -> ReportRow:
    """Helper creating a ReportRow with configurable defaults."""
    defaults = {
        "ioc": "evil.com",
        "vt_score": "5/66",
        "vendor_status": "Detected",
    }
    defaults.update(overrides)
    return ReportRow(**defaults)


CLEAN_ROW = make_row(
    ioc="8.8.8.8",
    vt_score="0/75",
    vendor_status="Clean",
)


def test_header_and_rows(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    write_report(p, [make_row(), CLEAN_ROW], META)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Results", "Summary"]
    ws = wb[RESULTS_SHEET]
    assert [c.value for c in ws[1]] == [
        "IOC",
        "Sophos Status",
        "VT Score",
    ]
    assert [c.value for c in ws[2]] == [
        "evil.com",
        "Detected",
        "5/66",
    ]
    assert [c.value for c in ws[3]] == [
        "8.8.8.8",
        "Clean",
        "0/75",
    ]
    assert ws.max_row == 3
    wb.close()


def test_vendor_in_header(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    meta = ReportMeta(
        vendor="CrowdStrike",
        input_name="list.txt",
        generated_at="2026-10-25 10:00:00 UTC",
    )
    write_report(p, [make_row()], meta)
    wb = openpyxl.load_workbook(p)
    ws = wb[RESULTS_SHEET]
    assert ws["B1"].value == "CrowdStrike Status"
    wb.close()


def test_freeze_and_filter(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    write_report(p, [make_row(), CLEAN_ROW], META)
    wb = openpyxl.load_workbook(p)
    ws = wb[RESULTS_SHEET]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C3"
    wb.close()


def test_empty_report(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    write_report(p, [], META)
    wb = openpyxl.load_workbook(p)
    ws = wb[RESULTS_SHEET]
    assert ws.max_row == 1
    assert ws.auto_filter.ref == "A1:C1"
    wb.close()
    assert read_report_rows(p) == []


def test_summary_sheet(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    rows = [
        make_row(vendor_status="Detected"),
        make_row(vendor_status="Detected"),
        make_row(vendor_status="Clean"),
        make_row(
            vendor_status="Unknown", vt_score="Unknown", vt_malicious=None
        ),
        make_row(vendor_status="Invalid"),
    ]
    write_report(p, rows, META)
    wb = openpyxl.load_workbook(p)
    summary = wb[SUMMARY_SHEET]
    assert summary["B1"].value == "Sophos"
    assert summary["B2"].value == "list.txt"
    assert summary["B3"].value == "2026-10-25 10:00:00 UTC"
    assert summary["B4"].value == 5
    assert summary["B5"].value == 2
    assert summary["A7"].value == "Status"
    assert summary["B7"].value == "Count"
    a_vals = [summary[f"A{i}"].value for i in range(8, 14)]
    b_vals = [summary[f"B{i}"].value for i in range(8, 14)]
    assert a_vals == [
        "Detected",
        "Clean",
        "Unknown",
        "Invalid",
        "Error",
        "Not checked",
    ]
    assert b_vals == [2, 1, 1, 1, 0, 0]
    wb.close()


def test_unsupported_status_rejected(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    with pytest.raises(ValueError, match="Unsupported status label"):
        write_report(p, [make_row(vendor_status="Maybe")], META)
    assert not p.exists()


def test_formula_injection_safe(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    rows = [
        make_row(ioc="=1+1", vendor_status="Invalid"),
        make_row(ioc="+cmd", vendor_status="Invalid"),
        make_row(ioc="-2+3", vendor_status="Invalid"),
        make_row(ioc="@SUM(A1)", vendor_status="Invalid"),
        make_row(vt_score='=HYPERLINK("http://x")', vendor_status="Invalid"),
    ]
    write_report(p, rows, META)
    wb = openpyxl.load_workbook(p)
    ws = wb[RESULTS_SHEET]
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=1+1"
    assert ws["A3"].data_type == "s" and ws["A3"].value == "+cmd"
    assert ws["A4"].data_type == "s" and ws["A4"].value == "-2+3"
    assert ws["A5"].data_type == "s" and ws["A5"].value == "@SUM(A1)"
    assert (
        ws["C6"].data_type == "s"
        and ws["C6"].value == '=HYPERLINK("http://x")'
    )
    wb.close()

    with zipfile.ZipFile(p, "r") as z:
        sheet1_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "<f>" not in sheet1_xml


def test_control_chars_and_length(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    rows = [
        make_row(ioc="evil\x0b.com"),
        make_row(ioc="a" * 40000),
    ]
    write_report(p, rows, META)
    read_back = read_report_rows(p)
    assert read_back[0].ioc == "evil.com"
    assert len(read_back[1].ioc) == 32767


def test_roundtrip(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    rows = [
        make_row(),
        CLEAN_ROW,
        make_row(
            ioc="0" * 64,
            vt_score="Unknown",
            vendor_status="Unknown",
        ),
        make_row(
            ioc="bad.example",
            vt_score="",
            vendor_status="Error",
        ),
        make_row(
            ioc="=1+1",
            vt_score="",
            vendor_status="Invalid",
        ),
    ]
    write_report(p, rows, META)
    back = read_report_rows(p)
    assert back == rows
    p.unlink()


def test_read_rejects_other_workbook(tmp_path: Path) -> None:
    p = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    wb.save(p)
    wb.close()
    with pytest.raises(ValueError, match="not a vt-ioc-checker report"):
        read_report_rows(p)


def test_read_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        read_report_rows(tmp_path / "nope.xlsx")


def test_status_fills_and_bold_header(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    rows = [
        make_row(vendor_status="Detected"),
        make_row(vendor_status="Clean"),
        make_row(vendor_status="Unknown"),
    ]
    write_report(p, rows, META)
    wb = openpyxl.load_workbook(p)
    ws = wb[RESULTS_SHEET]
    assert ws["A1"].font.bold is True
    cell_detected = ws["B2"]
    cell_clean = ws["B3"]
    cell_unknown = ws["B4"]
    assert cell_detected.fill.fill_type == "solid"
    assert str(cell_detected.fill.fgColor.rgb)[-6:] == "F8CBAD"
    assert cell_clean.fill.fill_type == "solid"
    assert str(cell_clean.fill.fgColor.rgb)[-6:] == "C6EFCE"
    assert cell_unknown.fill.fill_type == "solid"
    assert str(cell_unknown.fill.fgColor.rgb)[-6:] == "FFEB9C"
    wb.close()


def test_extension_check(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match=r"\.xlsx"):
        write_report(tmp_path / "r.csv", [], META)


def test_creates_parent_dir_no_leftover(tmp_path: Path) -> None:
    path = tmp_path / "sub" / "dir" / "r.xlsx"
    write_report(path, [], META)
    assert sorted(os.listdir(path.parent)) == ["r.xlsx"]


def test_overwrite_existing(tmp_path: Path) -> None:
    p = tmp_path / "r.xlsx"
    write_report(p, [make_row(), CLEAN_ROW], META)
    write_report(p, [make_row()], META)
    wb = openpyxl.load_workbook(p)
    ws = wb[RESULTS_SHEET]
    assert ws.max_row == 2
    wb.close()


def test_permission_error_message(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    p = tmp_path / "r.xlsx"

    def fake_replace(src, dst):
        raise PermissionError("Access denied")

    monkeypatch.setattr("vt_ioc_checker.report.os.replace", fake_replace)
    with pytest.raises(ReportWriteError, match="close it"):
        write_report(p, [make_row()], META)
    assert not p.exists()
    assert list(tmp_path.glob("*.partial*")) == []
